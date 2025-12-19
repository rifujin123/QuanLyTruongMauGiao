import io

from flask import Blueprint, send_file
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

from app.models.Models import get_monthly_and_yearly_revenue

report_br = Blueprint('report', __name__)

@report_br.route("/export-revenue")
def export_revenue_excel():
    monthly_rows, year_total_map = get_monthly_and_yearly_revenue()

    wb = Workbook()
    ws = wb.active
    ws.title = "Doanh thu"

    # Header
    ws["A1"] = "Tháng"
    ws["B1"] = "Năm"
    ws["C1"] = "Tổng doanh thu tháng"
    ws["D1"] = "Tổng doanh thu 12 tháng"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ["A1", "B1", "C1", "D1"]:
        ws[cell].fill = header_fill
        ws[cell].font = header_font
        ws[cell].alignment = Alignment(horizontal="center", vertical="center")

    # Data
    row_idx = 2
    for year, month, total in monthly_rows:
        ws[f"A{row_idx}"] = month
        ws[f"B{row_idx}"] = year
        ws[f"C{row_idx}"] = float(total or 0)

        year_total = float(year_total_map.get(year, 0))
        ws[f"D{row_idx}"] = year_total

        ws[f"C{row_idx}"].number_format = "#,##0"
        ws[f"D{row_idx}"].number_format = "#,##0"
        row_idx += 1

    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 24

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="bao_cao_doanh_thu_thang_nam.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )